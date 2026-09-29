import io
import json
import unittest
from unittest.mock import patch
from docx import Document
from openpyxl import Workbook
import tempfile
from pathlib import Path

import app
from ad_metrics import audit_report, number


def report(text, **meta):
    return {'client_name': '測試品牌', 'company_name': '測試公司', 'report_month': '2026-09',
            'platforms': {'meta': text}, 'platform_metadata': {'meta': {'currency': 'TWD', 'conversion_type': 'purchase', **meta}},
            'period_confirmed': True}


def entry(text, **meta):
    return audit_report(report(text, **meta))['platforms'][0]


class MetricsTests(unittest.TestCase):
    def test_weighted_ratios_and_decimal_precision(self):
        e = entry('名稱,花費,曝光,點擊,成果,轉換價值\nA,100,1000,10,2,300\nB,900,3000,90,8,1700')
        self.assertEqual(e['errors'], [])
        self.assertEqual(e['metrics'], {'spend': '1000', 'impressions': '4000', 'clicks': '100', 'conversions': '10', 'revenue': '2000', 'ctr': '2.5', 'cpc': '10', 'cpm': '250', 'cpa': '100', 'roas': '2'})
        self.assertEqual(entry('花費：0.1\n曝光：3\n點擊：1')['metrics']['cpc'], '0.1')

    def test_summary_rows_excluded_and_mismatched_totals_visible(self):
        e = entry('名稱 | 花費 | 曝光 | 點擊\nA | 100 | 1000 | 10\nB | 200 | 2000 | 20\n總計：全部 | 999 | 3000 | 30')
        self.assertEqual(e['metrics']['spend'], '300')
        self.assertEqual(e['excluded_total_lines'], [4])
        self.assertTrue(any('篩選範圍' in w for w in e['warnings']))

    def test_only_single_summary_is_accepted(self):
        self.assertEqual(entry('名稱,花費,曝光\n總計：平台,100,1000')['metrics']['spend'], '100')
        self.assertTrue(entry('名稱,花費,曝光\n總計：平台,100,1000\n總計：帳戶,100,1000')['errors'])

    def test_missing_values_do_not_become_zero_or_partial_totals(self):
        e = entry('名稱,花費,曝光,點擊\nA,100,1000,10\nB,,2000,20')
        self.assertIsNone(e['metrics']['spend'])
        self.assertIsNone(e['metrics']['cpc'])
        self.assertEqual(e['metrics']['impressions'], '3000')

    def test_zero_denominators_are_unavailable(self):
        e = entry('花費：0\n曝光：0\n點擊：0\n成果：0\n轉換價值：0')
        self.assertEqual(e['metrics']['spend'], '0')
        self.assertTrue(all(e['metrics'][k] is None for k in ['ctr', 'cpc', 'cpm', 'cpa', 'roas']))

    def test_thousands_quoting_empty_cells_and_multiline_provenance(self):
        e = entry('名稱,花費,曝光,點擊\n"名稱,含逗號","1,200",3000,30\n"跨\n行名稱",300,1000,10')
        self.assertEqual(e['metrics']['spend'], '1500')
        self.assertEqual([r['line'] for r in e['rows']], [2, 3])
        self.assertEqual(entry('名稱 | 花費 | 曝光\nA | 142,478 | 100000')['metrics']['spend'], '142478')

    def test_currency_conflict_and_unknown_currency(self):
        self.assertTrue(entry('名稱,花費,曝光,幣別\nA,100,1000,USD')['errors'])
        e = entry('花費：100\n曝光：1000\n點擊：10', currency='')
        self.assertIsNone(e['metrics']['spend'])
        self.assertEqual(e['metrics']['ctr'], '1')
        e = entry('名称,花費金額 (TWD),曝光\nA,100,1000', currency='')
        self.assertEqual(e['currency'], 'TWD')
        self.assertEqual(e['metrics']['spend'], '100')

    def test_different_conversion_events_are_not_summed(self):
        e = entry('名稱,花費,曝光,成果,轉換價值,成果指標\nA,100,1000,10,1000,purchase\nB,200,2000,20,2000,message', conversion_type='')
        for field in ['conversions', 'revenue', 'cpa', 'roas']:
            self.assertIsNone(e['metrics'][field])
        self.assertEqual(e['metrics']['spend'], '300')
        self.assertTrue(entry('名稱,花費,曝光,成果指標\nA,100,1000,message')['errors'])

    def test_clicks_are_not_inferred_from_interactions(self):
        e = entry('名稱,費用,曝光,互動\nA,100,1000,30')
        self.assertIsNone(e['metrics']['clicks'])
        self.assertIsNone(e['metrics']['ctr'])

    def test_imported_ratios_are_checked_not_averaged(self):
        e = entry('名稱,花費,曝光,點擊,CTR,CPC\nA,100,1000,10,80%,99\nB,900,3000,90,3%,10')
        self.assertEqual(e['metrics']['ctr'], '2.5')
        self.assertEqual(len(e['comparisons']), 2)
        self.assertEqual(e['comparisons'][0]['calculated'], '1')
        self.assertTrue(any('未標示百分比' in w for w in entry('花費：100\n曝光：1000\n點擊：10\nCTR：0.01')['warnings']))

    def test_bad_numbers_do_not_silently_parse(self):
        for value in ['1,23', '-1', 'NaN', 'Infinity', '1e4', '$100', '5%', '1.000,50', '1000000000000000000000']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                number(value)
        self.assertTrue(entry('花費：1,23\n曝光：1000')['errors'])
        self.assertTrue(entry('花費：100\n曝光：1.5')['errors'])

    def test_duplicate_rows_headers_and_malformed_rows_block(self):
        for text in [
            '名稱,花費,曝光\nA,100,1000\nA,100,1000',
            '名稱,花費,費用,曝光\nA,100,100,1000',
            '名稱,花費,曝光\nA,100,1000\n名稱,花費,曝光\nB,200,2000',
            '名稱,花費,曝光\nA,100,1000,extra',
            '花費：100\n花費：200\n曝光：1000',
        ]:
            with self.subTest(text=text):
                self.assertTrue(entry(text)['errors'])

    def test_period_mismatch_and_reverse_dates_block(self):
        for start, end in [('2026-08-01', '2026-09-30'), ('2026-09-30', '2026-09-01'), ('bad', '2026-09-30')]:
            self.assertTrue(entry(f'名稱,花費,曝光,報告開始,報告結束\nA,100,1000,{start},{end}')['errors'])

    def test_google_preamble_keeps_source_lines(self):
        e = entry('廣告活動報表\n2026-09-01 - 2026-09-30\n廣告活動,費用,曝光,點擊\nA,100,1000,10')
        self.assertEqual(e['rows'][0]['line'], 4)
        self.assertEqual(e['errors'], [])

    def test_no_cross_platform_conversion_or_roas_total(self):
        data = report('花費：100\n成果：2\n轉換價值：300')
        data['platforms']['google'] = '花費：200\n成果：2\n轉換價值：300'
        audit = audit_report(data)
        self.assertEqual(len(audit['platforms']), 2)
        self.assertNotIn('totals', audit)
        self.assertTrue(any('重複歸因' in w for w in audit['warnings']))

    def test_campaign_named_like_total_is_not_a_summary(self):
        e = entry('名稱,花費,曝光\n總計畫品牌,100,1000\n一般品牌,200,2000')
        self.assertEqual(e['metrics']['spend'], '300')
        self.assertEqual(e['excluded_total_lines'], [])

    def test_unrecognized_currency_is_not_overridden_by_metadata(self):
        self.assertTrue(entry('名稱,花費,曝光,幣別\nA,100,1000,UNKNOWN')['errors'])

    def test_unknown_dimension_keeps_distinct_rows(self):
        e = entry('campaign_id,花費,曝光\nID1,100,1000\nID2,100,1000')
        self.assertEqual(e['errors'], [])
        self.assertEqual(e['metrics']['spend'], '200')

    def test_invalid_input_boundaries(self):
        for data in [[], {}, report(''), report('x' * 200001), {**report('花費：100'), 'report_month': '2026-13'}, {**report('花費：100'), 'platform_metadata': []}]:
            with self.subTest(data=str(data)[:80]), self.assertRaises(ValueError):
                audit_report(data)


class MetricsApiTests(unittest.TestCase):
    def setUp(self):
        self.client = app.app.test_client()
        self.data = report('花費：100\n曝光：1000\n點擊：20\n成果：2\n轉換價值：300')
        self.model = {'overview': '需人工覆核', 'platforms': [{'index': 0, 'analysis': '先核對來源'}], 'insights': '資料有限', 'recommendations': '補充素材資料'}

    def approve(self):
        response = self.client.post('/api/ad-report/validate', json=self.data)
        self.assertEqual(response.status_code, 200)
        self.data['audit_fingerprint'] = response.json['fingerprint']

    def test_validate_does_not_use_ai(self):
        with patch.object(app, 'call_deepseek') as call:
            self.approve()
            call.assert_not_called()

    def test_unreviewed_stale_or_unconfirmed_input_never_calls_ai(self):
        with patch.object(app, 'call_deepseek') as call:
            self.assertEqual(self.client.post('/api/ad-report/process', json=self.data).status_code, 400)
            self.approve()
            self.data['platforms']['meta'] += '\nCPC：5'
            self.assertEqual(self.client.post('/api/ad-report/process', json=self.data).status_code, 400)
            self.data['period_confirmed'] = False
            self.approve()
            self.assertEqual(self.client.post('/api/ad-report/process', json=self.data).status_code, 400)
            call.assert_not_called()

    def test_report_metrics_are_server_owned_and_exported(self):
        self.approve()
        with patch.object(app, 'call_deepseek', return_value=json.dumps({**self.model, 'cover': {'title': '模型不該改標題'}, 'audit': {'metrics': 'fake'}})) as call:
            response = self.client.post('/api/ad-report/process', json=self.data)
        self.assertEqual(response.status_code, 200)
        result = response.json
        self.assertEqual(result['cover']['title'], '廣告月報')
        self.assertEqual(result['audit']['platforms'][0]['metrics']['ctr'], '2')
        self.assertIn('| CTR (%) | 2 |', result['sections'][1]['content'])
        self.assertIn('AI 解讀（需人工覆核）', result['sections'][1]['content'])
        self.assertNotIn('CPC：', str(call.call_args))
        self.assertNotIn('source_lines', str(call.call_args))
        self.assertNotIn('測試品牌', str(call.call_args))
        response = self.client.post('/api/ad-report/export', json=result)
        self.assertEqual(response.status_code, 200)
        doc = Document(io.BytesIO(response.data))
        self.assertTrue(any('CTR (%)' in c.text for t in doc.tables for r in t.rows for c in r.cells))
        self.assertTrue(any('點擊 ÷ 曝光 × 100' in p.text for p in doc.paragraphs))

    def test_model_schema_validation(self):
        self.approve()
        for payload in [[], {}, {**self.model, 'platforms': []}, {**self.model, 'platforms': [{'index': True, 'analysis': 'x'}]}, {**self.model, 'overview': 100}]:
            with patch.object(app, 'call_deepseek', return_value=json.dumps(payload)):
                self.assertEqual(self.client.post('/api/ad-report/process', json=self.data).status_code, 500)

    def test_validate_reports_unsupported_data_without_ai(self):
        self.data['platforms']['meta'] = '這次表現不錯，請幫我算'
        response = self.client.post('/api/ad-report/validate', json=self.data)
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json['valid'])

    def test_xlsx_preserves_embedded_commas_and_empty_cells(self):
        with tempfile.TemporaryDirectory() as folder:
            filename = Path(folder) / 'data.xlsx'
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(['名稱', '花費', '曝光', '點擊'])
            sheet.append(['品牌,台北', '1,200', 10000, 200])
            workbook.save(filename)
            parsed = entry(app.read_text_file(str(filename)))
            self.assertEqual(parsed['metrics']['spend'], '1200')
            self.assertEqual(parsed['rows'][0]['name'], '品牌,台北')
            self.assertEqual(parsed['errors'], [])


if __name__ == '__main__':
    unittest.main()
